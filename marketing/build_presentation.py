from pathlib import Path
import base64, html, json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'sunum'
OUT.mkdir(exist_ok=True)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(16), Inches(9)
NAVY, BLUE, TEAL, WHITE, MUTED = '0B1224','2864FF','43DFC0','F7F9FF','A9B7D0'
pages=[]
def slide(n, tag, dark=True):
    global s, els
    s=prs.slides.add_slide(prs.slide_layouts[6]); els=[]
    s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(NAVY)
    rect(60,49,8,29,TEAL)
    text(85,48,600,30,'RİTİM / QUALITY',23,WHITE,True)
    text(1050,53,490,25,tag,15,MUTED)
    rect(60,840,1480,1,'29344A')
    text(60,859,1200,20,'RİTİM AİLESİ  •  ÜRÜN TANITIMI',12,MUTED)
    text(1470,855,70,24,f'0{n} / 03',14,MUTED)
def rect(x,y,w,h,color):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x/100), Inches(y/100), Inches(w/100), Inches(h/100))
    sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(color); sh.line.fill.background()
    els.append(f'<div style="left:{x}px;top:{y}px;width:{w}px;height:{h}px;background:#{color}"></div>')
def text(x,y,w,h,value,size=24,color=WHITE,bold=False):
    sh=s.shapes.add_textbox(Inches(x/100), Inches(y/100), Inches(w/100), Inches(h/100))
    tf=sh.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
        p.font.name='Arial'; p.font.size=Pt(size*.72); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(0)
    els.append(f'<div style="left:{x}px;top:{y}px;width:{w}px;height:{h}px;font-size:{size}px;color:#{color};font-weight:{700 if bold else 400};line-height:1.15">{html.escape(value).replace(chr(10),"<br>")}</div>')
def pic(name,x,y,w,h,crop=(0,0,0,0)):
    path=ROOT/'screenshots'/name
    sh=s.shapes.add_picture(str(path), Inches(x/100), Inches(y/100), Inches(w/100), Inches(h/100))
    l,t,r,b=crop; sh.crop_left=l; sh.crop_top=t; sh.crop_right=r; sh.crop_bottom=b
    uri='data:image/png;base64,'+base64.b64encode(path.read_bytes()).decode()
    els.append(f'<div style="left:{x}px;top:{y}px;width:{w}px;height:{h}px;overflow:hidden"><img src="{uri}" style="position:absolute;width:{100/(1-l-r)}%;height:{100/(1-t-b)}%;left:{-100*l/(1-l-r)}%;top:{-100*t/(1-t-b)}%"></div>')
def done(notes):
    s.notes_slide.notes_text_frame.text=notes
    pages.append('<section class="slide">'+''.join(els)+'</section>')

slide(1,'01 — RİTİM QUALITY NEDİR?')
text(60,139,700,50,'Üretimin kalite hafızası.',48,WHITE,True)
text(60,218,580,145,'Her ölçümde\naynı standart.\nHer kararda veri.',54,WHITE,True)
text(60,434,535,112,'Kontrol planından operatör ölçümüne,\nölçüm kaydından süreç analizine:\nüretim kalitesi için tek dijital platform.',26,MUTED)
rect(60,617,492,56,BLUE)
text(80,632,465,30,'PLANLA  →  ÖLÇ  →  İZLE  →  İYİLEŞTİR',21,WHITE,True)
text(60,725,530,60,'Ritim ailesinin üretim kalitesine\nodaklanan çözümü.',23,TEAL)
rect(635,226,905,569,'202C44')
pic('01-olcum-istasyonu.png',647,238,881,545,(.18,.075,0,0))
text(647,194,850,25,'ÜRÜNDEN GÖRÜNÜM / ÖLÇÜM İSTASYONU',14,TEAL,True)
done('Ritim Quality, üretimde neyin nasıl kontrol edileceğini, ölçüm sonuçlarını ve bu sonuçlardan çıkan analizleri aynı akışta bir araya getirir. Operatör teknik resimdeki ölçüm noktasını takip eder; kalite ekibi planı ve toleransları yönetir. Ekran görüntüleri proje içindeki demo görselleridir; müşteri referansı veya performans taahhüdü değildir.')

slide(2,'02 — HANGİ SORUNLARI ÇÖZER?')
text(60,128,1460,67,'Dağınık kayıtlardan ortak bir kalite akışına.',46,WHITE,True)
cards=[(60,'01','Farklı dosyalar, farklı standartlar','Kontrol planlarını merkezileştirir.','Ölçüm noktaları, toleranslar ve\nrevizyonlar aynı yerde yönetilir.'),
       (565,'02','Operatöre bağlı uygulama farkları','Ölçümü adım adım yönlendirir.','Teknik resimde aktif nokta gösterilir;\ngirilen değer toleransla karşılaştırılır.'),
       (1070,'03','Geçmiş kayda ulaşma güçlüğü','Ölçümün bağlamını korur.','Ürün, lot, operatör ve zaman bilgisiyle\ngeçmiş sonuçlar takip edilir.')]
for x,num,problem,title,body in cards:
    rect(x,234,470,267,'142139'); text(x+24,257,75,40,num,31,TEAL,True)
    text(x+24,312,424,35,problem,20,MUTED)
    text(x+24,361,424,62,title,27,WHITE,True)
    text(x+24,436,424,56,body,20,MUTED)
pic('02-kontrol-planlari.png',60,533,880,253,(.20,.215,.025,.428))
text(977,554,550,40,'Standart tanımla. Güvenle uygula.',27,TEAL,True)
text(977,615,535,116,'Plan ve revizyon bilgisi ölçümle birlikte\nizlenir. Ekipler aynı kontrol düzeni\nüzerinden çalışır.',26,WHITE)
text(60,800,880,20,'ÜRÜNDEN GÖRÜNÜM / KONTROL PLANLARI · DEMO VERİ',12,MUTED)
done('Burada üç pratik soruna odaklanıyoruz: farklı dosyalarda tutulan kontrol bilgileri, operatörler arasında değişebilen uygulamalar ve geriye dönük kayıt arama zorluğu. Ritim Quality bunları merkezi kontrol planları, yönlendirilmiş ölçüm ve bağlamı korunan kayıtlarla ele alır. Amaç, kalite sürecini kişisel alışkanlıklara daha az bağımlı ve ekipler arasında daha tutarlı hale getirmektir.')

slide(3,'03 — RİTİM AİLESİNE NE KAZANDIRIR?')
text(60,128,1450,65,'Ölçümü kayda, veriyi karara dönüştürür.',48,WHITE,True)
text(60,208,1440,45,'Kalite ekibi için görünürlük. Operatör için netlik. Yönetim için ortak veri.',27,MUTED)
rect(60,293,885,449,'202C44')
pic('03-spc-analizi.png',72,305,861,425,(.205,.318,.03,0))
text(60,758,885,24,'ÜRÜNDEN GÖRÜNÜM / SPC ANALİZİ · DEMO VERİ',13,MUTED)
for y,num,title,body in [(299,'01','Süreç değişimini görün','Cp/Cpk, histogram ve kontrol grafikleriyle\nsüreç performansını değerlendirin.'),(456,'02','Geçmişi daha kolay inceleyin','Ölçüm kayıtları ve plan revizyonlarıyla\ninceleme ve raporlamayı destekleyin.'),(613,'03','İyileştirmeyi veriye dayandırın','Kalite ve üretim ekiplerini aynı ölçüm\nverisi etrafında buluşturun.')]:
    text(1000,y,55,40,num,27,TEAL,True); text(1066,y,474,43,title,28,WHITE,True); text(1066,y+56,474,72,body,23,MUTED)
text(60,794,1480,33,'Ritim Quality  /  Daha tutarlı kontroller. Daha izlenebilir üretim.',25,TEAL,True)
done('Ölçüm biriktirmek tek başına yeterli değildir. SPC, yani istatistiksel süreç kontrolü, ölçümlerdeki dağılımı ve zaman içindeki değişimi incelemeyi sağlar. Cp ve Cpk süreç yeteneğini değerlendirmede kullanılan göstergelerdir; sonuçlar kullanılan veriye ve analiz varsayımlarına bağlıdır. Ritim ailesi için değer önerimiz, ortak bir kalite dili ve veriye dayanan iyileştirme kararlarıdır. Kapanışta ölçüm istasyonu ve SPC ekranı üzerinden kısa bir canlı demo yapılabilir.')

prs.save(OUT/'Ritim-Quality-Tanitim.pptx')
css='''*{box-sizing:border-box}body{margin:0;background:#070C17;font-family:Arial,sans-serif}.slide{position:relative;width:1600px;height:900px;background:#0B1224;overflow:hidden;page-break-after:always}.slide>div{position:absolute}@page{size:16in 9in;margin:0}@media print{.slide{width:16in;height:9in;zoom:.96;break-inside:avoid}}@media screen{.slide{margin:25px auto}}'''
(OUT/'Ritim-Quality-Tanitim.html').write_text('<!doctype html><html lang="tr"><meta charset="utf-8"><title>Ritim Quality • Tanıtım</title><style>'+css+'</style>'+''.join(pages)+'</html>',encoding='utf-8')
(OUT/'Sunucu-Notlari.md').write_text('# Ritim Quality — Sunucu notları\n\n'+ '\n\n'.join(f'## Slayt {i+1}\n\n'+sl.notes_slide.notes_text_frame.text for i,sl in enumerate(prs.slides))+'\n\nGörsel kaynakları: marketing/screenshots klasöründeki mevcut üç uygulama ekran görüntüsü. Sunum metinleri mevcut ürün koduyla karşılaştırılmıştır.\n',encoding='utf-8')
print(str(OUT))
